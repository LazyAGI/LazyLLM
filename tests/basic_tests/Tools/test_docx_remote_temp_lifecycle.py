from io import BytesIO
import tempfile
import uuid

import docx
import fsspec
import pytest

from lazyllm.module.module import ModuleExecutionError
from lazyllm.tools.rag.readers.docxReader import DocxReader


def document_bytes():
    document = docx.Document()
    document.add_paragraph('native paragraph')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'native cell'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize('valid', [True, False])
def test_remote_reader_removes_owned_temp_on_success_and_failure(tmp_path, monkeypatch, valid):
    temp = tmp_path / 'temp'
    temp.mkdir()
    monkeypatch.setattr(tempfile, 'tempdir', str(temp))
    fs = fsspec.filesystem('memory')
    fs.auto_mkdir = True
    path = '/docx-temp-' + uuid.uuid4().hex + '.docx'
    fs.pipe(path, document_bytes() if valid else b'not a docx archive')
    try:
        reader = DocxReader(split_doc=True, save_image=False)
        if valid:
            nodes = reader(path, fs=fs)
        else:
            with pytest.raises(ModuleExecutionError):
                reader(path, fs=fs)
        assert list(temp.iterdir()) == []
        if valid:
            assert any('native cell' in node.get_text() for node in nodes)
            assert any('native paragraph' in node.get_text() for node in nodes)
    finally:
        fs.rm(path)


def test_local_rich_reader_keeps_existing_output(tmp_path):
    path = tmp_path / 'local.docx'
    path.write_bytes(document_bytes())
    nodes = DocxReader(split_doc=True, save_image=False)(path)
    assert any('native cell' in node.get_text() for node in nodes)
